#!/usr/bin/env python
"""Turn the Swahili briefs into one workbook that can be sent to the author.

A raw CSV dump is not sendable: 18 columns, notes up to 1,233 characters, and no
indication of which four columns the author is meant to type in. This adds only
presentation — no cell's VALUE is changed, which `verify()` asserts by reading
the workbook back and comparing it to the CSV.

Both sheets go in ONE workbook on purpose. The relations sheet is unusable
without the brief (it supplies the words every {REL} substitutes) and the brief
is unfinishable without the relations sheet, so they travel as one attachment.
Pass --split for two files instead.

    python review/csv_to_xlsx.py                 # -> ~/kinyamed-briefs/
    python review/csv_to_xlsx.py --out DIR --split

Needs openpyxl, which is not in requirements.txt because nothing else uses it:

    python3 -m venv /tmp/xlsxenv && /tmp/xlsxenv/bin/pip install openpyxl
    /tmp/xlsxenv/bin/python review/csv_to_xlsx.py
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
BRIEF = ROOT / "review" / "speaker_brief_swahili_v2.csv"
RELATIONS = ROOT / "review" / "speaker_brief_swahili_v2_relations.csv"

# Columns the author writes in. Tinted so that "where do I type" needs no
# explanation, and listed here rather than inferred so a new column does not
# silently join them.
AUTHOR_COLUMNS = {"your_phrasing", "second_phrasing_optional", "regional_variant",
                  "your_notes", "your_swahili"}

WIDTHS = {
    "concept_id": 11, "domain": 19, "proposed_urgency": 11, "english_gloss": 34,
    "person": 8, "person_note": 32, "applies": 8, "action": 24,
    "relation_set": 22, "relation_set_members": 28, "keep_distinct_from": 32,
    "hold": 7, "needs_clinician": 13, "brief_notes": 68,
    "english": 20, "used_by": 30, "note": 60,
}
DEFAULT_WIDTH = 30

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
AUTHOR_FILL = PatternFill("solid", fgColor="FFF6DC")
SKIP_FILL = PatternFill("solid", fgColor="F2F2F2")


def write_sheet(ws, path: Path, freeze: str) -> None:
    rows = list(csv.DictReader(path.open(encoding="utf-8")))
    columns = list(rows[0])

    for i, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(name, DEFAULT_WIDTH)
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(rows, start=2):
        # A row the author is told to skip is greyed rather than hidden. Hiding
        # it would lose the reason it is skipped, which is the part worth
        # reading — several of them record where a collapsed concept went.
        skipped = row.get("action", "").startswith("SKIP")
        for c, name in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=row[name])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if name in AUTHOR_COLUMNS:
                cell.fill = AUTHOR_FILL
            elif skipped:
                cell.fill = SKIP_FILL
                cell.font = Font(color="808080")

    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def start_here(ws) -> None:
    """The rules that must survive the file being forwarded on its own.

    Deliberately short. `docs/swahili-authoring-brief.md` is the full page and
    this does not try to replace it — it exists so that a workbook opened
    without the covering email is still usable.
    """
    lines = [
        ("Writing the Swahili phrases", True),
        ("", False),
        ("Two tabs: 'brief' is the 128 concepts, 'relations' is the 12 relation "
         "words that every third-person phrase needs.", False),
        ("The full instructions are in swahili-authoring-brief.md, sent with this "
         "file.", False),
        ("", False),
        ("WHERE TO TYPE", True),
        ("The cream-coloured columns are yours. Everything else is ours - please "
         "leave it as it is.", False),
        ("Grey rows are ones you are asked to SKIP; the reason is in brief_notes.", False),
        ("", False),
        ("WHAT WE ARE ASKING FOR", True),
        ("For each row marked WRITE, put in your_phrasing the sentence a patient "
         "would actually say in Kiswahili for the meaning in english_gloss.", False),
        ("This is not a translation task. The gloss is a clinical description, "
         "deliberately not written as a sentence, so there is nothing to "
         "translate and nothing to edit.", False),
        ("There is no Swahili anywhere in this workbook, on purpose. A draft to "
         "correct pulls corrections towards the draft, and its mistakes survive. "
         "You are the first person to write these.", False),
        ("", False),
        ("THE FIVE RULES", True),
        ("1. Write a complete sentence in the patient's voice. Capital letter, "
         "full stop.", False),
        ("2. No time reference inside it - the generator adds one ('since "
         "yesterday'). Unless the duration IS the concept, as in a cough lasting "
         "more than two weeks.", False),
        ("3. Do not end with an added clause - the generator adds one of those too.",
         False),
        ("4. Third-person rows: write {REL} where the relation word goes, those "
         "five characters exactly. Make it the grammatical subject, and check the "
         "sentence works for every relation in relation_set_members.", False),
        ("5. Never mix first and third person inside one sentence.", False),
        ("", False),
        ("THREE QUESTIONS", True),
        ("CR05 (wheeze), GI03 and GI05 (a word for stool), PA08 (a child's ear "
         "discharging) are stuck in both other languages. The full question is in "
         "brief_notes on those rows. 'There is no ordinary word for this' is a "
         "real answer and a useful one.", False),
    ]
    ws.column_dimensions["A"].width = 110
    for r, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = Font(bold=bold, size=13 if bold and r == 1 else 11)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def verify(path: Path, sheets: list[tuple[str, Path]]) -> None:
    """Read the workbook back and assert every cell equals the CSV it came from.

    The conversion is presentation only. Anything else - a dropped row, a value
    coerced to a number or a date, a truncated note - is a corrupted brief being
    sent to a speaker, and it would not be visible by looking at the file.
    """
    book = load_workbook(path)
    for sheet_name, source in sheets:
        rows = list(csv.DictReader(source.open(encoding="utf-8")))
        columns = list(rows[0])
        ws = book[sheet_name]
        if ws.max_row != len(rows) + 1:
            raise SystemExit(f"{sheet_name}: {ws.max_row - 1} rows, CSV has {len(rows)}")
        for r, row in enumerate(rows, start=2):
            for c, name in enumerate(columns, start=1):
                got = ws.cell(row=r, column=c).value
                got = "" if got is None else str(got)
                if got != row[name]:
                    raise SystemExit(
                        f"{sheet_name} row {r} {name}: workbook has {got!r}, "
                        f"CSV has {row[name]!r}")
    print(f"  verified {path.name}: every cell matches its CSV")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=Path.home() / "kinyamed-briefs")
    ap.add_argument("--split", action="store_true",
                    help="Two workbooks instead of one with both tabs.")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    if args.split:
        for stem, source, freeze in [("speaker_brief_swahili_v2", BRIEF, "B2"),
                                     ("speaker_brief_swahili_v2_relations",
                                      RELATIONS, "A2")]:
            book = Workbook()
            write_sheet(book.active, source, freeze)
            book.active.title = "brief" if source is BRIEF else "relations"
            path = args.out / f"{stem}.xlsx"
            book.save(path)
            verify(path, [(book.active.title, source)])
            print(f"wrote {path}")
        return 0

    book = Workbook()
    start_here(book.active)
    book.active.title = "start here"
    write_sheet(book.create_sheet("brief"), BRIEF, "B2")
    write_sheet(book.create_sheet("relations"), RELATIONS, "A2")
    path = args.out / "speaker_brief_swahili_v2.xlsx"
    book.save(path)
    verify(path, [("brief", BRIEF), ("relations", RELATIONS)])
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
